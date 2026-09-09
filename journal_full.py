"""journal_full.py — ONE journal: every trade TAKEN and every alert MISSED.

The daily journal.xlsx is trades-only. This merges the taken trades (master_ledger.csv
via ledger.py — the truth source since 9/9, not the truncating day table) with
the misses (misses.py, from trades.log)
into a SINGLE sheet you can filter and mine to tune the app: which callers, which
reasons, which contracts got skipped and why — right next to what actually filled
and how it did.

    python3 journal_full.py              every day  -> journal-full-ALL.xlsx
    python3 journal_full.py --today      today only -> journal-full-<date>.xlsx
    python3 journal_full.py --date 2026-09-08

Read-only. Reads master_ledger.csv + trades.log; writes an .xlsx. Never trades.
"""
import os
import re
import sys
import datetime as _dt

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import ledger
import misses as _m

HERE = os.path.dirname(os.path.abspath(__file__))
DAYS = os.path.join(HERE, "days")

COLS = ["date", "time", "outcome", "reason", "room", "caller", "symbol", "side",
        "direction", "contract", "qty", "avg_in", "exits", "P&L", "P&L %",
        "max run-up %", "max drawdown %", "opened", "closed", "status",
        "exit_by", "account", "signal"]


# Discord alerts can carry ANSI colour codes and other control bytes; openpyxl
# refuses those. Strip them so a stray escape can't blow up the whole export.
_ANSI = re.compile(r"\x1b\[[0-9;]*m")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(v):
    if isinstance(v, str):
        v = _ANSI.sub("", v)
        v = _ILLEGAL.sub("", v)
    return v


def _hhmm(t):
    if not t:
        return ""
    try:
        return _dt.datetime.fromtimestamp(t).strftime("%H:%M")
    except Exception:                                       # noqa: BLE001
        return ""


def _taken_rows(date=None, all_days=False):
    rows = []
    # 9/9: reads master_ledger.csv via ledger.py — days/*.json table truncates.
    for dd, day_rows in sorted(ledger.by_day().items()):
        if not all_days and dd != date:
            continue
        for r in day_rows:
            is_call = str(r.get("side") or "").upper().startswith("C")
            ct = ""
            side_col = ""
            direction = ""
            if r.get("strike") is not None:
                ct = "%s%s %s" % (r.get("strike"), "C" if is_call else "P",
                                  r.get("expiry") or "")
                side_col = "CALL" if is_call else "PUT"
            elif r.get("kind") == "future":
                ct = "futures"
                direction = "SHORT" if int(r.get("direction") or 1) < 0 else "LONG"
            ex = "; ".join(
                "%s@%s%s" % (e.get("qty"), e.get("price"),
                             "" if e.get("pl") is None else " (%+.0f)" % e["pl"])
                for e in (r.get("exits") or []))
            pl_pct = r.get("pl_pct")
            hi = r.get("hi_pct")
            lo = r.get("lo_pct")
            rows.append({
                "date": dd, "time": _hhmm(r.get("opened")),
                "outcome": "TAKEN", "reason": "",
                "room": r.get("room") or "", "caller": r.get("who") or "?",
                "symbol": r.get("symbol") or "", "side": side_col,
                "direction": direction, "contract": ct, "qty": r.get("qty") or 0,
                "avg_in": r.get("avg") if r.get("avg") is not None else "",
                "exits": ex, "P&L": r.get("pl"),
                "P&L %": ("%+.1f%%" % pl_pct) if pl_pct is not None else "",
                "max run-up %": ("%+.1f%%" % hi) if hi is not None else "",
                "max drawdown %": ("%+.1f%%" % lo) if lo is not None else "",
                "opened": _hhmm(r.get("opened")), "closed": _hhmm(r.get("closed")),
                "status": r.get("state") or "", "exit_by": r.get("exit_by") or "",
                "account": r.get("account") or ("live" if r.get("live") else "paper"),
                "signal": r.get("raw") or "",
            })
    return rows


def _missed_rows(date=None, all_days=False):
    rows = []
    for m in _m.collect_misses(date=date, all_days=all_days):
        rows.append({
            "date": m["date"], "time": m["time"], "outcome": "MISSED",
            "reason": m["reason"], "room": "", "caller": "",
            "symbol": m["symbol"] if m["symbol"] != "?" else "",
            "side": "", "direction": "", "contract": "", "qty": "",
            "avg_in": "", "exits": "", "P&L": "", "P&L %": "",
            "max run-up %": "", "max drawdown %": "", "opened": "",
            "closed": "", "status": "", "exit_by": "", "account": "",
            "signal": m["raw"],
        })
    return rows


def _count(rows, key):
    d = {}
    for r in rows:
        k = r.get(key) or "?"
        d[k] = d.get(k, 0) + 1
    return sorted(d.items(), key=lambda x: -x[1])


def main(argv):
    date = None
    all_days = True
    i = 0
    while i < len(argv):
        a = argv[i]
        if a == "--today":
            date = _dt.date.today().isoformat()
            all_days = False
        elif a == "--all":
            all_days = True
            date = None
        elif a == "--date" and i + 1 < len(argv):
            date = argv[i + 1]
            all_days = False
            i += 1
        i += 1

    taken = _taken_rows(date=date, all_days=all_days)
    missed = _missed_rows(date=date, all_days=all_days)
    allrows = taken + missed
    allrows.sort(key=lambda r: (r["date"], r["time"]))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All trades"
    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")
    ws.append(COLS)
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(vertical="center")
    take_fill = PatternFill("solid", fgColor="ECFDF5")   # green tint = taken
    miss_fill = PatternFill("solid", fgColor="FEF2F2")   # red tint = missed
    for r in allrows:
        ws.append([_clean(r.get(k, "")) for k in COLS])
        fill = take_fill if r["outcome"] == "TAKEN" else miss_fill
        for c in ws[ws.max_row]:
            c.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"date": 11, "time": 6, "outcome": 8, "reason": 24, "room": 16,
              "caller": 16, "symbol": 8, "side": 6, "direction": 9,
              "contract": 16, "exits": 24, "signal": 64}
    for idx, k in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(k, 10)

    sm = wb.create_sheet("Summary")
    sm.append(["UNIFIED JOURNAL — " + ("ALL DAYS" if all_days else date)])
    sm.append([])
    sm.append(["Taken", len(taken)])
    sm.append(["Missed", len(missed)])
    sm.append(["Total", len(allrows)])
    sm.append([])
    sm.append(["MISSED by reason", ""])
    for reason, n in _count(missed, "reason"):
        sm.append([reason, n])
    sm.append([])
    sm.append(["MISSED by symbol (top 20)", ""])
    for sym, n in _count([r for r in missed if r["symbol"]], "symbol")[:20]:
        sm.append([sym, n])
    sm.append([])
    sm.append(["TAKEN by caller", ""])
    for who, n in _count(taken, "caller"):
        sm.append([who, n])
    sm["A1"].font = Font(bold=True, size=13)
    for row in sm.iter_rows():
        if row[0].value in ("MISSED by reason", "MISSED by symbol (top 20)",
                            "TAKEN by caller"):
            row[0].font = Font(bold=True)
    sm.column_dimensions["A"].width = 30
    sm.column_dimensions["B"].width = 10

    scope = "ALL" if all_days else date
    out = os.path.join(HERE, "journal-full-%s.xlsx" % scope)
    wb.save(out)
    print("wrote %s" % out)
    print("  taken=%d  missed=%d  total=%d"
          % (len(taken), len(missed), len(allrows)))


if __name__ == "__main__":
    main(sys.argv[1:])
